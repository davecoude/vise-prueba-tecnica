"""Lector: obtiene las cédulas crudas del Excel de entrada. No valida su formato (eso es del validador)."""

import logging
import unicodedata
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import FilaCedula

logger = logging.getLogger(__name__)

COLUMNA = "cedula"


class ArchivoEntradaError(Exception):
    """El archivo no existe, no es un Excel válido o no tiene la columna `cedula`."""


def _normalizar_encabezado(valor: object) -> str:
    # "Cédula ", "CEDULA" y "cedula" se aceptan como la misma columna
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return texto.strip().lower()


def leer_cedulas(ruta: Path) -> list[FilaCedula]:
    if not ruta.is_file():
        raise ArchivoEntradaError(f"No existe el archivo de entrada: {ruta}")

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as error:
        raise ArchivoEntradaError(f"No se pudo abrir {ruta} como archivo Excel (.xlsx): {error}") from error

    try:
        hoja = libro.active
        if hoja is None:
            raise ArchivoEntradaError(f"El archivo {ruta} no tiene hojas")

        filas = hoja.iter_rows(values_only=True)
        encabezados = next(filas, None)
        if not encabezados:
            raise ArchivoEntradaError(f"El archivo {ruta} está vacío")

        nombres = [_normalizar_encabezado(h) for h in encabezados]
        if COLUMNA not in nombres:
            raise ArchivoEntradaError(
                f"El archivo {ruta} no tiene la columna '{COLUMNA}'. Columnas encontradas: {list(encabezados)}"
            )
        indice = nombres.index(COLUMNA)

        resultado: list[FilaCedula] = []
        for numero, fila in enumerate(filas, start=2):  # la fila 1 es el encabezado
            if fila is None or all(celda is None for celda in fila):
                continue  # filas totalmente vacías (artefacto común de Excel)
            valor = fila[indice] if indice < len(fila) else None
            resultado.append(FilaCedula(fila=numero, valor=valor))
    finally:
        libro.close()

    logger.info("Archivo leído: %s | %d filas con datos", ruta, len(resultado))
    return resultado
