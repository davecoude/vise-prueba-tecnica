"""Generador del reporte Excel y del resumen en consola."""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Estado, ResultadoConsulta, Resumen

logger = logging.getLogger(__name__)

COLUMNAS = ["cedula", "nombres", "estado", "detalle", "fecha_consulta"]
ANCHOS = [14, 34, 18, 70, 20]

VERDE_VISE = "1B5E20"
COLOR_ESTADO = {
    Estado.LIMPIO: "E7F3E8",
    Estado.CON_ANTECEDENTES: "FBEFD9",
    Estado.NO_ENCONTRADO: "E6EDF7",
    Estado.ERROR: "FBE9EA",
}


def escribir_reporte(resultados: list[ResultadoConsulta], resumen: Resumen, ruta: Path) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro = Workbook()

    hoja = libro.active
    assert hoja is not None
    hoja.title = "Reporte"
    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE_VISE)
        celda.alignment = Alignment(vertical="center")

    for r in resultados:
        # La cédula se escribe como texto: conserva ceros a la izquierda y evita notación científica
        hoja.append([r.cedula, r.nombres, r.estado.value, r.detalle, r.fecha_consulta.replace(tzinfo=None)])
        fila = hoja.max_row
        hoja.cell(fila, 1).number_format = "@"
        hoja.cell(fila, 3).fill = PatternFill("solid", fgColor=COLOR_ESTADO[r.estado])
        hoja.cell(fila, 5).number_format = "yyyy-mm-dd hh:mm:ss"

    for indice, ancho in enumerate(ANCHOS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    hoja_resumen = libro.create_sheet("Resumen")
    for etiqueta, valor in _filas_resumen(resumen):
        hoja_resumen.append([etiqueta, valor])
    hoja_resumen.column_dimensions["A"].width = 32
    hoja_resumen.column_dimensions["B"].width = 22
    for celda in hoja_resumen["A"]:
        celda.font = Font(bold=True)

    libro.save(ruta)
    logger.info("Reporte generado: %s (%d filas)", ruta, len(resultados))
    return ruta


def _filas_resumen(resumen: Resumen) -> list[tuple[str, object]]:
    return [
        ("Fecha de ejecución", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total en el reporte", resumen.total),
        ("Limpios", resumen.limpios),
        ("Con antecedentes", resumen.con_antecedentes),
        ("No encontrados", resumen.no_encontrados),
        ("Errores", resumen.errores),
        ("  de ellos, cédulas inválidas", resumen.invalidas),
        ("Duplicados eliminados", resumen.duplicados_eliminados),
    ]


def formatear_resumen(resumen: Resumen) -> str:
    """Resumen para consola (total, limpios, con antecedentes, errores + contexto)."""
    lineas = [
        "",
        "=" * 46,
        " RESUMEN DE LA CONSULTA DE ANTECEDENTES",
        "=" * 46,
        f" Total consultado/reportado : {resumen.total:>6}",
        f" Limpios                    : {resumen.limpios:>6}",
        f" Con antecedentes           : {resumen.con_antecedentes:>6}",
        f" No encontrados             : {resumen.no_encontrados:>6}",
        f" Errores                    : {resumen.errores:>6}",
        f"   (cédulas inválidas)      : {resumen.invalidas:>6}",
        f" Duplicados eliminados      : {resumen.duplicados_eliminados:>6}",
        "=" * 46,
    ]
    return "\n".join(lineas)
