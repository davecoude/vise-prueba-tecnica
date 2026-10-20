import socket
from datetime import datetime

from openpyxl import load_workbook

from rpa_antecedentes.client import HttpAntecedentesClient
from rpa_antecedentes.models import Estado, ResultadoConsulta
from rpa_antecedentes.pipeline import ejecutar
from rpa_antecedentes.report import COLUMNAS, formatear_resumen

FECHA = datetime(2026, 10, 20, 9, 30)


class ClienteFalso:
    """Test double: respuestas predefinidas por cédula y registro de llamadas."""

    def __init__(self, respuestas: dict[str, Estado | Exception]) -> None:
        self.respuestas = respuestas
        self.consultadas: list[str] = []

    def consultar(self, cedula: str) -> ResultadoConsulta:
        self.consultadas.append(cedula)
        respuesta = self.respuestas.get(cedula, Estado.LIMPIO)
        if isinstance(respuesta, Exception):
            raise respuesta
        return ResultadoConsulta(cedula, "Nombre Ficticio", respuesta, "detalle", FECHA)


def leer_reporte(ruta):
    hoja = load_workbook(ruta)["Reporte"]
    filas = list(hoja.iter_rows(values_only=True))
    return filas[0], [dict(zip(filas[0], fila, strict=True)) for fila in filas[1:]]


def test_flujo_completo_con_datos_mixtos(crear_excel, tmp_path) -> None:
    entrada = crear_excel(
        [["1032456781"], ["12345"], ["79845612"], ["1032456781"], ["ABC1234567"], ["52987431"]]
    )
    salida = tmp_path / "out" / "reporte_antecedentes.xlsx"
    cliente = ClienteFalso({"79845612": Estado.CON_ANTECEDENTES, "52987431": Estado.NO_ENCONTRADO})

    resumen = ejecutar(entrada, salida, cliente, reloj=lambda: FECHA)

    # Solo se consultan las válidas, una vez cada una
    assert cliente.consultadas == ["1032456781", "79845612", "52987431"]
    assert (resumen.total, resumen.limpios, resumen.con_antecedentes, resumen.no_encontrados) == (5, 1, 1, 1)
    assert (resumen.errores, resumen.invalidas, resumen.duplicados_eliminados) == (2, 2, 1)

    encabezados, filas = leer_reporte(salida)
    assert list(encabezados) == COLUMNAS
    assert [f["cedula"] for f in filas] == ["1032456781", "12345", "79845612", "ABC1234567", "52987431"]
    assert filas[1]["estado"] == "error" and "Cédula inválida (fila 3)" in filas[1]["detalle"]
    assert filas[2]["estado"] == "con antecedentes"
    assert isinstance(filas[0]["fecha_consulta"], datetime)


def test_una_consulta_que_lanza_excepcion_no_detiene_el_proceso(crear_excel, tmp_path) -> None:
    entrada = crear_excel([["1032456781"], ["79845612"], ["52987431"]])
    cliente = ClienteFalso({"79845612": RuntimeError("fallo inesperado")})

    resumen = ejecutar(entrada, tmp_path / "r.xlsx", cliente, reloj=lambda: FECHA)

    assert cliente.consultadas == ["1032456781", "79845612", "52987431"]
    assert (resumen.limpios, resumen.errores) == (2, 1)


def test_con_el_servicio_caido_se_genera_el_reporte_completo_con_errores(crear_excel, tmp_path) -> None:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        puerto = s.getsockname()[1]
    cliente = HttpAntecedentesClient(f"http://127.0.0.1:{puerto}", timeout_s=1, max_reintentos=1, sleep=lambda _: None)
    entrada = crear_excel([["1032456781"], ["79845612"], ["123"]])
    salida = tmp_path / "reporte.xlsx"

    resumen = ejecutar(entrada, salida, cliente)

    assert salida.exists()
    assert resumen.total == 3 and resumen.errores == 3
    _, filas = leer_reporte(salida)
    assert all(f["estado"] == "error" for f in filas)
    assert "no se pudo conectar" in filas[0]["detalle"]


def test_archivo_sin_cedulas_genera_reporte_vacio(crear_excel, tmp_path) -> None:
    resumen = ejecutar(crear_excel([]), tmp_path / "r.xlsx", ClienteFalso({}))

    assert resumen.total == 0


def test_resumen_de_consola_incluye_los_totales_exigidos(crear_excel, tmp_path) -> None:
    resumen = ejecutar(crear_excel([["1032456781"]]), tmp_path / "r.xlsx", ClienteFalso({}))

    texto = formatear_resumen(resumen)

    for etiqueta in ("Total", "Limpios", "Con antecedentes", "Errores"):
        assert etiqueta in texto
